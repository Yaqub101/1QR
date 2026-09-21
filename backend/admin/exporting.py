"""CSV and XLSX output of a Report, and the audit row every export leaves behind (SYSTEM_SPEC 20).

CSV is UTF-8 WITH a byte-order mark: that is what makes Excel open Devanagari, accented and CJK names
correctly instead of as mojibake. Both formats guard against spreadsheet formula injection: a text cell that
starts with = + - @ (a name or a reason typed by a person) must never run as a formula when opened.
"""
from __future__ import annotations

import csv
import io
import re

from openpyxl import Workbook
from openpyxl.styles import Font
from sqlalchemy.engine import Connection

from backend.admin.reports import Report
from backend.audit import write_audit

CSV_TYPE = "text/csv; charset=utf-8"
XLSX_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
_FORMULA_START = ("=", "+", "-", "@", "\t", "\r")
_ILLEGAL_XML = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def _clean(value):
    return _ILLEGAL_XML.sub("", value) if isinstance(value, str) else value


def _csv_cell(value):
    value = _clean(value)
    if isinstance(value, str) and value.startswith(_FORMULA_START):
        return "'" + value  # the standard neutraliser: Excel shows the text, does not evaluate it
    return value


def to_csv(report: Report) -> bytes:
    out = io.StringIO(newline="")
    writer = csv.writer(out)
    writer.writerow([label for _, label in report.columns])
    for row in report.rows:
        writer.writerow([_csv_cell(row.get(key, "")) for key, _ in report.columns])
    return out.getvalue().encode("utf-8-sig")


def to_xlsx(report: Report) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = report.key[:31]
    ws.append([label for _, label in report.columns])
    for row in report.rows:
        ws.append([_clean(row.get(key, "")) for key, _ in report.columns])
    for cells in ws.iter_rows():
        for cell in cells:
            if isinstance(cell.value, str):
                cell.data_type = "s"  # a text cell stays text, even when it starts with "="
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "A2"
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


FORMATS = {"csv": (to_csv, CSV_TYPE), "xlsx": (to_xlsx, XLSX_TYPE)}


def log_export(conn: Connection, principal, report: Report, fmt: str, params: dict) -> None:
    """Who exported what. Written in the same transaction that read the data: no log row, no file."""
    write_audit(conn, "EXPORT", operator_id=principal.user_id,
                details={"report": report.key, "format": fmt, "rows": len(report.rows), "filters": {k: v for k, v in params.items() if v}})
