"""Phases 13 + 16: the Admin dashboard and every report, against ONE constructed dataset with a known answer.

HOW THE ANSWERS ARE KNOWN (so a test can fail):
  * The dataset below is built row by row as raw SQL, and each student's expected journey is written down in
    Python at the same moment (`Person.active`). Every expected number is computed from that Python model, and
    the headline ones are also written out by hand as literals (see test_hand_calculated_headline_numbers).
  * The independent SQL in this file deliberately uses a DIFFERENT formulation from the code under test: it links
    a reversal to its original through corrects_event_id, while backend/admin matches on (student, activity,
    completion_cycle). Two definitions that agree on a dataset with reversals in it are evidence, not tautology.

Like the other Phase 7-12 suites: one shared PostgreSQL test database, three venue app instances (plus central),
no sync.
"""
import io
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook
from sqlalchemy import text

from backend.stage import state as stage_state
from tests.admin_support import ORDER, S1, S2, S3, add_event, add_student, build_dataset, parse_csv, rows, scalar
from tests.test_auth import ACTIVITIES
from tests.test_station_engine import (  # noqa: F401  (engine / world / apps are pytest fixtures)
    _CLIENTS,
    admin,
    apps,
    confirm,
    engine,
    make_student,
    operator,
    world,
)



@pytest.fixture(scope="module", autouse=True)
def _fresh_client_cache(world):
    _CLIENTS.clear()
    yield
    _CLIENTS.clear()


# --------------------------------------------------------------------------- the dataset
@pytest.fixture(scope="module")
def data(engine, world):
    return build_dataset(engine)


def model(data, activity):
    return {p.prn for p in data.all if activity in p.active}


# --------------------------------------------------------------------------- independent SQL
ACTIVE_STUDENTS = """
    SELECT DISTINCT student_id FROM activity_events
    WHERE activity = :a AND kind IN ('COMPLETE','WAIVER')
      AND event_id NOT IN (SELECT corrects_event_id FROM activity_events WHERE kind = 'REVERSAL')"""


def sql_active(engine, activity):
    return {r["prn"] for r in rows(engine, f"SELECT prn FROM students WHERE id IN ({ACTIVE_STUDENTS})", a=activity)}


def get(apps, venue, path, **kw):
    return admin(apps, venue).get(path, **kw)


def dash(apps, venue="hall"):
    r = get(apps, venue, "/admin/api/dashboard")
    assert r.status_code == 200, r.text
    return r.json()


# =========================================================================== DASHBOARD COUNTS
class TestDashboardCounts:
    def test_hand_calculated_headline_numbers(self, apps, data):
        """Worked out on paper from the dataset above, not from any query."""
        c = dash(apps)["counts"]
        assert (c["registered"], c["reported"], c["yet_to_report"], c["not_attended"], c["registration_reversed"]) == (22, 17, 5, 4, 1)
        assert c["reporting_percent"] == 77.3  # 17 / 22
        funnel = {f["activity"]: f["count"] for f in dash(apps)["funnel"]}
        assert funnel == {"REGISTRATION": 17, "THOBE_ALLOCATION": 14, "SEATING": 12, "QUEUE": 10,
                          "STAGE": 7, "THOBE_RETURN": 4, "LUNCH": 2}
        assert dash(apps)["outstanding_thobes"]["count"] == 10

    def test_every_count_matches_a_direct_database_query(self, apps, engine, data):
        d = dash(apps)
        c = d["counts"]
        assert c["registered"] == scalar(engine, "SELECT count(*) FROM students")
        assert c["reported"] == len(sql_active(engine, "REGISTRATION"))
        assert c["not_attended"] == scalar(
            engine, "SELECT count(*) FROM students WHERE id NOT IN (SELECT student_id FROM activity_events WHERE activity = 'REGISTRATION')")
        assert c["yet_to_report"] == scalar(engine, "SELECT count(*) FROM students") - len(sql_active(engine, "REGISTRATION"))
        assert c["reporting_percent"] == round(100 * len(sql_active(engine, "REGISTRATION")) / c["registered"], 1)
        for step in d["funnel"]:  # the funnel: one independent query per activity
            assert step["count"] == len(sql_active(engine, step["activity"])), step["activity"]
        assert d["exceptions"]["open"] == scalar(engine, "SELECT count(*) FROM exceptions WHERE status = 'OPEN'")
        assert d["exceptions"]["provisional_events"] == scalar(engine, "SELECT count(*) FROM activity_events WHERE flags @> ARRAY['PROVISIONAL']")
        assert d["exceptions"]["manual_entries"] == scalar(engine, "SELECT count(*) FROM activity_events WHERE flags @> ARRAY['MANUAL']")
        assert d["exceptions"]["duplicate_attempts"] == scalar(engine, "SELECT count(*) FROM scan_log WHERE result = 'DUPLICATE'")
        assert d["exceptions"]["blocked_attempts"] == scalar(engine, "SELECT count(*) FROM scan_log WHERE result = 'REJECTED'")
        assert d["exceptions"]["corrections"] == scalar(engine, "SELECT count(*) FROM activity_events WHERE kind IN ('REVERSAL','WAIVER')")

    def test_registered_equals_reported_plus_yet_to_report(self, apps):
        c = dash(apps)["counts"]
        assert c["registered"] == c["reported"] + c["yet_to_report"]
        assert c["yet_to_report"] == c["not_attended"] + c["registration_reversed"]  # every non-reporter is accounted for

    def test_the_model_agrees_with_the_database_for_every_activity(self, engine, data):
        """If this fails the DATASET is wrong (not the app): it ties the Python model to the raw rows."""
        for activity in ORDER:
            assert model(data, activity) == sql_active(engine, activity), activity

    def test_school_wise_reporting(self, apps, engine, data):
        by_school = {r["school"]: r for r in dash(apps)["school_wise"]}
        assert set(by_school) == {S1, S2, S3}
        for school in (S1, S2, S3):
            members = [p for p in data.all if p.school == school]
            got = by_school[school]
            assert got["registered"] == len(members) == scalar(engine, "SELECT count(*) FROM students WHERE school = :s", s=school)
            assert got["reported"] == sum(1 for p in members if "REGISTRATION" in p.active)
            assert got["registered"] == got["reported"] + got["yet_to_report"]
        assert [by_school[s]["registered"] for s in (S1, S2, S3)] == [10, 6, 6]   # by hand
        assert [by_school[s]["reported"] for s in (S1, S2, S3)] == [8, 4, 5]     # by hand
        assert sum(r["reported"] for r in by_school.values()) == dash(apps)["counts"]["reported"]

    def test_funnel_marks_the_waived_returns(self, apps):
        returned = next(f for f in dash(apps)["funnel"] if f["activity"] == "THOBE_RETURN")
        assert returned["count"] == 4 and returned["of_which_waived"] == 1

    def test_outstanding_thobes_are_exactly_allocated_and_not_returned(self, apps, engine, data):
        out = dash(apps)["outstanding_thobes"]
        expected = {p.prn for p in data.all if "THOBE_ALLOCATION" in p.active and "THOBE_RETURN" not in p.active}
        via_sql = sql_active(engine, "THOBE_ALLOCATION") - sql_active(engine, "THOBE_RETURN")
        assert expected == via_sql and len(expected) == out["count"] == 10
        assert {s["prn"] for s in out["students"]} == expected
        assert data.people["L"].prn in expected      # a reversed return is outstanding again
        assert data.people["J"].prn not in expected  # a waiver is not outstanding
        assert data.people["I"].prn not in expected and data.people["K1"].prn not in expected

    def test_stage_view_shows_current_led_and_the_waiting_queue(self, apps, data):
        st = dash(apps)["stage"]
        assert st["current"]["name"] == data.people["F2"].name
        assert st["led_mode"] == "SHOWING" and st["led_name"] == "Frank Two-Display"
        assert st["waiting"] == 1 and [n["name"] for n in st["next"]] == [data.people["F1"].name]

    def test_exception_counters_and_breakdown(self, apps):
        e = dash(apps)["exceptions"]
        assert e["open"] == 2 and e["open_by_type"] == {"CONFLICT": 1, "RETURN_WAIVED": 1}  # the resolved SEQ_GAP is not counted
        assert (e["provisional_events"], e["manual_entries"], e["duplicate_attempts"], e["blocked_attempts"]) == (1, 2, 3, 2)
        assert e["corrections"] == 3  # B's reversal, J's waiver, L's reversal

    def test_server_health_reports_database_and_backup_status(self, apps):
        h = dash(apps)["health"]
        assert h["database"] == "up"
        assert "last_backup_at" in h

    def test_the_dashboard_pages_render_the_same_numbers(self, apps):
        page = get(apps, "hall", "/admin/dashboard")
        live = get(apps, "hall", "/admin/dashboard/live")
        assert page.status_code == live.status_code == 200
        for html in (page.text, live.text):
            assert '<b id="n-registered">22</b>' in html and '<b id="n-reported">17</b>' in html
            assert '<b id="n-not-attended">4</b>' in html and 'id="n-outstanding">10<' in html
        assert "<html" in page.text and "<html" not in live.text  # the refresh returns only the swappable body


# =========================================================================== REPORTS
def report(apps, key, **params):
    r = get(apps, "hall", f"/admin/api/reports/{key}", params=params)
    assert r.status_code == 200, (key, r.text)
    return r.json()


def slug(activity):
    return "activity-" + activity.lower().replace("_", "-")


class TestPerActivityReports:
    @pytest.mark.parametrize("activity", ORDER)
    def test_total_equals_completed_plus_not_completed(self, apps, engine, data, activity):
        body = report(apps, slug(activity))
        totals, listed = body["totals"], body["rows"]
        expected_done = model(data, activity)
        # against the constructed dataset (known answer) ...
        assert totals["total"] == len(data.all) == scalar(engine, "SELECT count(*) FROM students")
        assert totals["completed"] == len(expected_done) and totals["not_completed"] == len(data.all) - len(expected_done)
        assert totals["total"] == totals["completed"] + totals["not_completed"]
        # ... and the listed rows really are the two halves: everyone exactly once, in the right half
        assert len(listed) == totals["total"] and len({r["prn"] for r in listed}) == totals["total"]
        assert {r["prn"] for r in listed if r["state"] == "Completed"} == expected_done == sql_active(engine, activity)
        assert {r["prn"] for r in listed if r["state"] == "Not Completed"} == {p.prn for p in data.all} - expected_done

    @pytest.mark.parametrize("activity", ORDER)
    def test_the_status_filter_returns_one_half_but_keeps_the_full_totals(self, apps, data, activity):
        done = report(apps, slug(activity), status="completed")
        rest = report(apps, slug(activity), status="not_completed")
        assert {r["prn"] for r in done["rows"]} == model(data, activity)
        assert len(done["rows"]) + len(rest["rows"]) == done["totals"]["total"] == rest["totals"]["total"]

    def test_a_not_completed_row_says_why(self, apps, data):
        reg = {r["prn"]: r for r in report(apps, slug("REGISTRATION"))["rows"]}
        assert reg[data.people["B"].prn]["note"] == "Reversed by Admin: wrong student scanned"
        stage = {r["prn"]: r for r in report(apps, slug("STAGE"))["rows"]}
        assert stage[data.people["H"].prn]["note"] == "Skipped: not ready"
        assert stage[data.people["G1"].prn]["state"] == "Completed" and stage[data.people["G1"].prn]["note"] == ""  # skipped, then completed

    def test_a_waiver_counts_as_a_completed_thobe_return(self, apps, data):
        rows_ = {r["prn"]: r for r in report(apps, slug("THOBE_RETURN"))["rows"]}
        assert rows_[data.people["J"].prn]["state"] == "Completed" and rows_[data.people["J"].prn]["record_kind"] == "WAIVER"
        assert rows_[data.people["L"].prn]["state"] == "Not Completed"  # reversed


class TestAttendanceAndJourneyReports:
    def test_not_attended_is_exactly_the_students_with_no_registration_event(self, apps, engine, data):
        body = report(apps, "not-attended")
        expected = {p.prn for p in data.all if not p.registration_event}
        via_sql = {r["prn"] for r in rows(engine, "SELECT prn FROM students WHERE id NOT IN "
                                                   "(SELECT student_id FROM activity_events WHERE activity = 'REGISTRATION')")}
        got = [r["prn"] for r in body["rows"]]
        assert set(got) == expected == via_sql and len(got) == len(set(got)) == body["totals"]["not_attended"] == 4
        assert data.people["B"].prn not in got  # registered once (then reversed): has an event, so not "never reported"
        assert all(data.people[k].prn in got for k in ("A1", "A2", "A3", "A4"))

    def test_not_attended_plus_reversed_plus_reported_reconciles_to_the_master_count(self, apps):
        c = dash(apps)["counts"]
        assert c["not_attended"] + c["registration_reversed"] + c["reported"] == c["registered"]

    def test_incomplete_journey_is_registered_but_not_exited(self, apps, data):
        body = report(apps, "incomplete-journey")
        expected = {p.prn for p in data.all if "REGISTRATION" in p.active and "LUNCH" not in p.active}
        assert {r["prn"] for r in body["rows"]} == expected and body["totals"]["incomplete"] == len(expected) == 15
        by = {r["prn"]: r for r in body["rows"]}
        # Seating is optional, so it is never listed as not done (Phase R4)
        assert by[data.people["C3"].prn]["not_yet_done"].startswith("Robe Allocation, Queue, Stage")
        assert by[data.people["L"].prn]["journey_status"] in ("Robe not returned", "Robe and money not returned")  # stage done, return reversed
        assert data.people["K1"].prn not in by and data.people["A1"].prn not in by and data.people["B"].prn not in by

    def test_stage_completed_and_skipped_with_reasons(self, apps, data):
        body = report(apps, "stage-outcomes")
        skipped = {r["prn"]: r["reason"] for r in body["rows"] if r["outcome"] == "Skipped"}
        assert skipped == {data.people["G1"].prn: "microphone problem", data.people["H"].prn: "not ready"}
        assert body["totals"] == {"completed": 7, "skipped": 2}


class TestThobeReports:
    def test_outstanding_thobes_list(self, apps, data):
        body = report(apps, "outstanding-robes")
        expected = {p.prn for p in data.all if "THOBE_ALLOCATION" in p.active and "THOBE_RETURN" not in p.active}
        assert {r["prn"] for r in body["rows"]} == expected and body["totals"]["outstanding"] == 10

    def test_waived_or_lost_list(self, apps, data):
        body = report(apps, "waived-robes")
        assert [r["prn"] for r in body["rows"]] == [data.people["J"].prn] and body["rows"][0]["reason"] == "lost robe"
        assert body["totals"] == {"waived": 1}

    def test_thobe_stock_check(self, apps):
        t = report(apps, "robe-count")["totals"]
        assert t == {"issued": 14, "returned": 3, "waived": 1, "outstanding": 10}  # 14 out; 3 back; 1 written off; 10 still out
        assert t["issued"] == t["returned"] + t["waived"] + t["outstanding"]


class TestFlagAndCorrectionReports:
    def test_late_provisional_and_manual_lists(self, apps, data):
        late = report(apps, "late-reporting")["rows"]
        assert {r["prn"] for r in late} == {data.people["C1"].prn, data.people["E1"].prn} and all(r["activity"] == "Reporting" for r in late)
        assert [r["prn"] for r in report(apps, "provisional")["rows"]] == [data.people["I"].prn]
        manual = report(apps, "manual")["rows"]
        assert {(r["prn"], r["activity"]) for r in manual} == {(data.people["C2"].prn, "Reporting"), (data.people["F2"].prn, "Queue")}

    def test_corrections_report_links_each_correction_to_its_original(self, apps, data):
        body = report(apps, "corrections")
        assert body["totals"] == {"corrections": 3}
        by = {r["prn"]: r for r in body["rows"]}
        assert by[data.people["B"].prn]["kind"] == "Reversal" and by[data.people["B"].prn]["reason"] == "wrong student scanned"
        assert by[data.people["B"].prn]["original_event_id"] == str(data.people["B"].events["REGISTRATION"])
        assert by[data.people["J"].prn]["kind"] == "Return waived / lost" and by[data.people["J"].prn]["original_event_id"] == ""
        assert by[data.people["L"].prn]["original_event_id"] == str(data.people["L"].events["THOBE_RETURN"])

    def test_exceptions_report_counts_open_and_resolved(self, apps):
        assert report(apps, "exceptions")["totals"] == {"open": 2, "resolved": 1}


class TestSummaries:
    def test_school_summary_matches_the_model_and_its_totals_reconcile_to_the_master_count(self, apps, data):
        body = report(apps, "school-summary")
        got = {r["school"]: r for r in body["rows"]}
        assert set(got) == {S1, S2, S3}
        names = {"reported": "REGISTRATION", "thobe_received": "THOBE_ALLOCATION",
                 "seated": "SEATING", "queued": "QUEUE", "stage_complete": "STAGE", "thobe_returned": "THOBE_RETURN",
                 "exited": "LUNCH"}
        for school in (S1, S2, S3):
            members = [p for p in data.all if p.school == school]
            assert got[school]["registered"] == len(members)
            assert got[school]["not_attended"] == sum(1 for p in members if not p.registration_event)
            for column, activity in names.items():
                assert got[school][column] == sum(1 for p in members if activity in p.active), (school, column)
        totals = body["totals"]
        assert totals["registered"] == len(data.all) == 22 and totals["reported"] == 17 and totals["not_attended"] == 4
        assert totals["registered"] == sum(r["registered"] for r in body["rows"])

    def test_programme_summary_splits_by_programme(self, apps, data):
        body = report(apps, "programme-summary")
        assert all("programme" in r for r in body["rows"]) and body["totals"]["registered"] == 22


# =========================================================================== EXPORTS
def export(apps, key, fmt="csv", venue="hall", **params):
    return get(apps, venue, f"/admin/api/reports/{key}/export", params={"format": fmt, **params})


class TestExports:
    def test_csv_round_trips_non_ascii_names_exactly(self, apps, data):
        r = export(apps, "not-attended")
        assert r.status_code == 200 and r.headers["content-type"].startswith("text/csv")
        assert "attachment" in r.headers["content-disposition"] and ".csv" in r.headers["content-disposition"]
        header, parsed = parse_csv(r.content)  # asserts the byte-order mark Excel needs
        assert header[:3] == ["PRN", "Name", "School"]
        names = {row["PRN"]: row["Name"] for row in parsed}
        assert names[data.people["A1"].prn] == "अनिल कुमार"   # Devanagari
        assert names[data.people["A2"].prn] == "José Müller"      # accents
        assert names[data.people["A3"].prn] == "李小龍"            # CJK
        assert len(parsed) == 4

    def test_csv_and_json_report_the_same_rows(self, apps):
        for key in ("school-summary", "activity-lunch", "waived-robes"):
            _, parsed = parse_csv(export(apps, key).content)
            assert len(parsed) == len(report(apps, key)["rows"]), key

    def test_xlsx_round_trips_non_ascii_names_and_has_a_header(self, apps, data):
        r = export(apps, "not-attended", "xlsx")
        assert r.status_code == 200 and "spreadsheetml" in r.headers["content-type"] and ".xlsx" in r.headers["content-disposition"]
        sheet = load_workbook(io.BytesIO(r.content)).active
        table = list(sheet.iter_rows(values_only=True))
        assert list(table[0][:3]) == ["PRN", "Name", "School"] and len(table) == 5
        names = {row[0]: row[1] for row in table[1:]}
        assert names[data.people["A1"].prn] == "अनिल कुमार" and names[data.people["A3"].prn] == "李小龍"

    def test_a_name_that_looks_like_a_formula_never_runs_as_one(self, apps, data):
        prn = data.people["A4"].prn
        _, parsed = parse_csv(export(apps, "not-attended").content)
        cell = next(r["Name"] for r in parsed if r["PRN"] == prn)
        assert cell == "'=SUM(1+1)"  # neutralised: Excel shows the text, does not evaluate it
        sheet = load_workbook(io.BytesIO(export(apps, "not-attended", "xlsx").content)).active
        found = next(c for row in sheet.iter_rows() for c in row if c.value == "=SUM(1+1)")
        assert found.data_type == "s"  # stored as text, not as a formula

    def test_every_report_in_the_catalogue_runs_and_exports_in_both_formats(self, apps, data):
        catalogue = get(apps, "hall", "/admin/api/reports").json()["reports"]
        keys = {c["key"] for c in catalogue}
        assert {"not-attended", "incomplete-journey", "outstanding-robes", "waived-robes", "robe-count", "late-reporting",
                "provisional", "manual", "corrections", "exceptions", "school-summary", "programme-summary", "stage-outcomes",
                "audit", "student-history"} <= keys
        assert {slug(a) for a in ORDER} <= keys
        sid = str(data.people["J"].s.id)
        for key in keys:
            params = {"student_id": sid} if key == "student-history" else {}
            assert report(apps, key, **params)["columns"], key
            for fmt in ("csv", "xlsx"):
                r = export(apps, key, fmt, **params)
                assert r.status_code == 200 and len(r.content) > 20, (key, fmt)

    def test_student_history_export_is_that_students_full_journey(self, apps, data):
        _, parsed = parse_csv(export(apps, "student-history", student_id=str(data.people["L"].s.id)).content)
        assert [r["Activity"] for r in parsed].count("Robe Return") == 2  # the return and its reversal are both kept
        assert {r["State"] for r in parsed if r["Activity"] == "Robe Return"} == {"REVERSED", "CORRECTION"}

    def test_an_export_leaves_an_audit_row_naming_who_and_what(self, apps, engine, world):
        before = scalar(engine, "SELECT count(*) FROM audit_log WHERE action = 'EXPORT'")
        assert export(apps, "corrections").status_code == 200
        last = rows(engine, "SELECT operator_id, details FROM audit_log WHERE action = 'EXPORT' ORDER BY id DESC LIMIT 1")[0]
        assert scalar(engine, "SELECT count(*) FROM audit_log WHERE action = 'EXPORT'") == before + 1
        assert last["operator_id"] == world.admin_id and last["details"]["report"] == "corrections"
        assert last["details"]["format"] == "csv" and last["details"]["rows"] == 3

    def test_bad_requests_are_plain_errors_not_crashes(self, apps):
        assert export(apps, "no-such-report").status_code == 404
        assert export(apps, "not-attended", "pdf").status_code == 400
        assert get(apps, "hall", "/admin/api/reports/student-history").status_code == 400
        assert get(apps, "hall", "/admin/api/reports/activity-seating", params={"status": "bogus"}).status_code == 400

    def test_the_report_pages_render(self, apps):
        assert get(apps, "hall", "/admin/reports").status_code == 200
        page = get(apps, "hall", "/admin/reports/not-attended")
        assert page.status_code == 200 and "अनिल कुमार" in page.text and "Download CSV" in page.text


# =========================================================================== LIVE UPDATE (must stay LAST)
class TestLiveUpdate:
    """Adds a student, so every other test in this module (which asserts exact totals) must run before it."""

    def test_dashboard_reflects_a_scan_made_elsewhere_immediately(self, apps, engine, world, data):
        before = dash(apps)["counts"]
        s = make_student(engine)
        reporting = operator(apps, world, "REGISTRATION")
        assert confirm(reporting, "REGISTRATION", token=s.token).json()["result"] == "CONFIRMED"
        after = dash(apps)["counts"]  # no cache, no wait: the very next read
        assert after["registered"] == before["registered"] + 1 and after["reported"] == before["reported"] + 1
        assert after["yet_to_report"] == before["yet_to_report"] and after["not_attended"] == before["not_attended"]
        assert after["reported"] == len(sql_active(engine, "REGISTRATION"))
